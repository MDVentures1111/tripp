# Synthetic, non-confidential fixtures for D04-C policy tests.
import io, zipfile, pikepdf, os
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document

def proposal_pdf(path, form=False):
    c = canvas.Canvas(path, pagesize=letter)
    c.setTitle("Synthetic Landlord Proposal")
    lines = ["SYNTHETIC LANDLORD PROPOSAL - TEST FIXTURE",
             "Premises: Suite 400, 100 Example Plaza, Testville",
             "Rentable Square Feet: 12,345", "Base Rent: $32.50/RSF/yr, 3% annual escalations",
             "Term: 84 months", "Free Rent: 6 months", "TI Allowance: $55.00/RSF"]
    for i, l in enumerate(lines):
        c.drawString(72, 720 - 18*i, l)
    if form:
        c.acroForm.textfield(name="tenant_signature", x=72, y=500, width=200, height=20)
    c.showPage(); c.drawString(72, 720, "Page 2 - synthetic terms"); c.save()

proposal_pdf("/tmp/base.pdf")
with pikepdf.open("/tmp/base.pdf") as p:
    p.save("proposal-objstm.pdf", object_stream_mode=pikepdf.ObjectStreamMode.generate, compress_streams=True)
proposal_pdf("/tmp/form.pdf", form=True)
with pikepdf.open("/tmp/form.pdf") as p:
    p.save("proposal-acroform.pdf", object_stream_mode=pikepdf.ObjectStreamMode.generate)
with pikepdf.open("/tmp/base.pdf") as p:
    p.Root.OpenAction = pikepdf.Array([p.pages[0].obj, pikepdf.Name.Fit])
    p.save("proposal-openaction.pdf")
with pikepdf.open("/tmp/base.pdf") as p:
    # A stream whose (uncompressed) payload contains marker-like bytes must not trigger rejection.
    s = pikepdf.Stream(p, b"q /JS 1 Tf /Launch BT ET Q")
    p.pages[0].Contents = pikepdf.Array([p.pages[0].Contents, p.make_indirect(s)]) if not isinstance(p.pages[0].Contents, pikepdf.Array) else p.pages[0].Contents
    p.pages[0].Contents.append(p.make_indirect(s))
    p.save("proposal-binary-markers.pdf", compress_streams=False)
with pikepdf.open("/tmp/base.pdf") as p:
    p.Root.OpenAction = pikepdf.Dictionary(S=pikepdf.Name.JavaScript, JS=pikepdf.String("app.alert('x')"))
    p.save("reject-javascript.pdf", object_stream_mode=pikepdf.ObjectStreamMode.disable)
with pikepdf.open("/tmp/base.pdf") as p:
    p.save("reject-encrypted.pdf", encryption=pikepdf.Encryption(owner="owner-test", user="user-test", R=4))

def docx_bytes():
    d = Document(); d.add_heading("Synthetic Letter of Intent - Test Fixture", 1)
    d.add_paragraph("Tenant: Example Tenant LLC. Premises: 9,876 RSF. Base rent $28.00/RSF.")
    t = d.add_table(rows=2, cols=2); t.cell(0,0).text = "Term"; t.cell(0,1).text = "60 months"
    b = io.BytesIO(); d.save(b); return b.getvalue()
raw = docx_bytes()
open("loi.docx", "wb").write(raw)

class NoSeek(io.RawIOBase):
    def __init__(self): self.buf = bytearray()
    def writable(self): return True
    def write(self, b): self.buf += b; return len(b)
src = zipfile.ZipFile(io.BytesIO(raw))
ns = NoSeek()
with zipfile.ZipFile(ns, "w", zipfile.ZIP_DEFLATED) as z:
    for info in src.infolist(): z.writestr(info.filename, src.read(info.filename))
open("loi-data-descriptor.docx", "wb").write(bytes(ns.buf))

def rewrite(extra_name, out):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist(): z.writestr(info.filename, src.read(info.filename))
        z.writestr(extra_name, b"synthetic-not-a-real-binary")
    open(out, "wb").write(b.getvalue())
rewrite("word/vbaProject.bin", "reject-macro.docx")
rewrite("word/embeddings/oleObject1.bin", "reject-ole.docx")
rewrite("word/embeddings/Microsoft_Excel_Worksheet.xlsx", "loi-embedded-xlsx.docx")
for f in sorted(os.listdir(".")): print(f, os.path.getsize(f))
